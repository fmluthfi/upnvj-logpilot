from __future__ import annotations

from io import BytesIO
import re
from zipfile import ZipFile

import pdfplumber
from docx import Document

from app.services.google_docs_parser import (
    ParsedLogbookRow,
    _field_for_header,
    parse_document,
    parse_indonesian_date,
)


MAX_UPLOAD_BYTES = 10 * 1024 * 1024
_MAX_DOCX_UNCOMPRESSED_BYTES = 50 * 1024 * 1024
_MAX_PDF_PAGES = 100
_DATE_START = re.compile(
    r"^\s*(?:(?:senin|selasa|rabu|kamis|jumat|jum'at|sabtu|minggu)\s*,?\s+)?"
    r"(?P<date>\d{4}-\d{1,2}-\d{1,2}|\d{1,2}[/.\-]\d{1,2}[/.\-]\d{4}|"
    r"\d{1,2}\s+[A-Za-z]+\s+\d{4})(?P<rest>.*)$",
    re.IGNORECASE,
)
_LABEL = re.compile(
    r"^(uraian(?:\s+kegiatan)?|kegiatan|aktivitas|activity|description|hasil|result|"
    r"kendala|hambatan|obstacle|dokumentasi|documentation)\s*:\s*(.*)$",
    re.IGNORECASE,
)


class FileLogbookError(ValueError):
    pass


def parse_uploaded_logbook(filename: str, content: bytes) -> list[ParsedLogbookRow]:
    if not content or len(content) > MAX_UPLOAD_BYTES:
        raise FileLogbookError("File harus berisi data dan tidak lebih dari 10 MB.")

    extension = filename.rsplit(".", 1)[-1].casefold() if "." in filename else ""
    try:
        if extension == "pdf" and content.startswith(b"%PDF-"):
            with pdfplumber.open(BytesIO(content)) as pdf:
                if len(pdf.pages) > _MAX_PDF_PAGES:
                    raise FileLogbookError("PDF melebihi batas 100 halaman.")
                tables = [table for page in pdf.pages for table in page.extract_tables()]
                text = "\n".join(page.extract_text() or "" for page in pdf.pages)
        elif extension == "docx" and content.startswith(b"PK"):
            with ZipFile(BytesIO(content)) as archive:
                if "word/document.xml" not in archive.namelist() or sum(
                    info.file_size for info in archive.infolist()
                ) > _MAX_DOCX_UNCOMPRESSED_BYTES:
                    raise FileLogbookError("File Word tidak valid atau terlalu besar setelah dibuka.")
            document = Document(BytesIO(content))
            tables = [[[cell.text for cell in row.cells] for row in table.rows] for table in document.tables]
            text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        else:
            raise FileLogbookError("Gunakan file PDF atau Word .docx yang valid. Format .doc lama belum didukung.")
    except FileLogbookError:
        raise
    except Exception as error:
        raise FileLogbookError("File tidak dapat dibaca. Periksa apakah file rusak atau dilindungi sandi.") from error

    rows = _parse_tables(tables)
    if not rows:
        rows = _parse_dated_text(text)
    if not rows:
        raise FileLogbookError(
            "Tidak ada catatan bertanggal yang terbaca. PDF hasil scan memerlukan OCR dan belum didukung."
        )
    return rows


def _parse_tables(tables: list[list[list[str | None]]]) -> list[ParsedLogbookRow]:
    groups: list[list[list[str]]] = []
    for table in tables:
        clean = [[str(cell or "").strip() for cell in row] for row in table if row]
        if not clean:
            continue
        if groups and not _has_header(clean) and len(clean[0]) == len(groups[-1][0]):
            groups[-1].extend(clean)
        else:
            groups.append(clean)

    document = {"body": {"content": [
        {"table": {"tableRows": [
            {"tableCells": [_cell(value) for value in row]}
            for row in _join_continuation_rows(group)
        ]}}
        for group in groups
    ]}}
    return parse_document(document)


def _has_header(rows: list[list[str]]) -> bool:
    return any(
        {"date", "activity"}.issubset({_field_for_header(value) for value in row})
        for row in rows[:6]
    )


def _join_continuation_rows(rows: list[list[str]]) -> list[list[str]]:
    header = next((row for row in rows[:6] if {"date", "activity"}.issubset(
        {_field_for_header(value) for value in row}
    )), None)
    if header is None:
        return rows
    date_column = next(i for i, value in enumerate(header) if _field_for_header(value) == "date")
    activity_column = next(i for i, value in enumerate(header) if _field_for_header(value) == "activity")
    joined: list[list[str]] = []
    for row in rows:
        if (joined and len(row) > max(date_column, activity_column)
            and not row[date_column] and row[activity_column]
            and date_column < len(joined[-1])
            and parse_indonesian_date(joined[-1][date_column])):
            for index, value in enumerate(row):
                if value and index < len(joined[-1]):
                    joined[-1][index] = "\n".join(filter(None, (joined[-1][index], value)))
            continue
        joined.append(row.copy())
    return joined


def _parse_dated_text(text: str) -> list[ParsedLogbookRow]:
    records: list[dict[str, str]] = []
    current: dict[str, str] | None = None
    field = "activity"
    for raw_line in text.splitlines():
        line = raw_line.strip()
        if not line:
            continue
        match = _DATE_START.match(line)
        if match and parse_indonesian_date(match.group("date")):
            current = {"date": match.group("date"), "activity": "", "result": "", "obstacle": "", "documentation": ""}
            records.append(current)
            field = "activity"
            line = match.group("rest").lstrip(" :–—-\t")
            if not line:
                continue
        if current is None:
            continue
        label = _LABEL.match(line)
        if label:
            name = label.group(1).casefold()
            field = ("result" if name in {"hasil", "result"} else
                     "obstacle" if name in {"kendala", "hambatan", "obstacle"} else
                     "documentation" if name in {"dokumentasi", "documentation"} else "activity")
            line = label.group(2)
        if line:
            current[field] = "\n".join(filter(None, (current[field], line)))

    if not records:
        return []
    table = [["Tanggal", "Uraian Kegiatan", "Hasil", "Kendala", "Dokumentasi"]]
    table.extend([[record[key] for key in ("date", "activity", "result", "obstacle", "documentation")]
                  for record in records])
    return _parse_tables([table])


def _cell(value: str) -> dict:
    return {"content": [{"paragraph": {"elements": [{"textRun": {"content": value + "\n"}}]}}]}
