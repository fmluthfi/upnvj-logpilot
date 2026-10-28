import unittest
from io import BytesIO
from unittest.mock import patch

from docx import Document

from app.services.file_logbook_parser import MAX_UPLOAD_BYTES, FileLogbookError, parse_uploaded_logbook


def word_bytes(paragraphs: list[str] = (), table_rows: list[list[str]] = ()) -> bytes:
    document = Document()
    for text in paragraphs:
        document.add_paragraph(text)
    if table_rows:
        table = document.add_table(rows=0, cols=len(table_rows[0]))
        for values in table_rows:
            cells = table.add_row().cells
            for cell, value in zip(cells, values):
                cell.text = value
    output = BytesIO()
    document.save(output)
    return output.getvalue()


class FakePage:
    def __init__(self, tables, text=""):
        self.tables = tables
        self.text = text

    def extract_tables(self):
        return self.tables

    def extract_text(self):
        return self.text


class FakePdf:
    def __init__(self, pages):
        self.pages = pages

    def __enter__(self):
        return self

    def __exit__(self, *_):
        return False


class FileLogbookParserTests(unittest.TestCase):
    def test_word_table_with_alternate_headers(self) -> None:
        content = word_bytes(table_rows=[
            ["Date", "Description", "Outcome"],
            ["21 Oktober 2026", "Menyusun rancangan", "Rancangan awal"],
        ])
        rows = parse_uploaded_logbook("catatan.docx", content)
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0].original_activity, "Menyusun rancangan")
        self.assertEqual(rows[0].result_text, "Rancangan awal")
        self.assertEqual(rows[0].kind, "filled")

    def test_word_without_table_reads_dated_paragraphs(self) -> None:
        content = word_bytes(paragraphs=[
            "Catatan magang",
            "Senin, 21 September 2026",
            "Kegiatan: Menyusun kebutuhan aplikasi",
            "Hasil: Daftar kebutuhan awal",
            "22 September 2026: Meninjau rancangan",
        ])
        rows = parse_uploaded_logbook("catatan.docx", content)
        self.assertEqual([row.original_activity for row in rows], [
            "Menyusun kebutuhan aplikasi", "Meninjau rancangan",
        ])
        self.assertEqual(rows[0].result_text, "Daftar kebutuhan awal")

    def test_pdf_joins_a_row_split_across_pages(self) -> None:
        pdf = FakePdf([
            FakePage([[["Tanggal", "Uraian Kegiatan"], ["21 September 2026", "Menyusun"]]]),
            FakePage([[["", "rancangan"]]]),
        ])
        with patch("app.services.file_logbook_parser.pdfplumber.open", return_value=pdf):
            rows = parse_uploaded_logbook("catatan.pdf", b"%PDF-1.4")
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0].original_activity, "Menyusun\nrancangan")

    def test_rejects_invalid_and_unreadable_files(self) -> None:
        for filename, content in (("catatan.doc", b"word"), ("catatan.pdf", b"not-pdf")):
            with self.subTest(filename=filename), self.assertRaises(FileLogbookError):
                parse_uploaded_logbook(filename, content)
        with patch("app.services.file_logbook_parser.pdfplumber.open", return_value=FakePdf([FakePage([])])):
            with self.assertRaisesRegex(FileLogbookError, "OCR"):
                parse_uploaded_logbook("scan.pdf", b"%PDF-1.4")

    def test_rejects_files_over_the_upload_limit(self) -> None:
        with self.assertRaisesRegex(FileLogbookError, "10 MB"):
            parse_uploaded_logbook("catatan.pdf", b"x" * (MAX_UPLOAD_BYTES + 1))


if __name__ == "__main__":
    unittest.main()
